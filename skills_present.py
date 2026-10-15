# -*- coding: utf-8 -*-
"""Навыки Ульвара: презентации PowerPoint (.pptx)."""
from skills import skill, _ok, _err, downloads_folder
import os
import re


def _parse_outline(outline):
    slides, cur = [], None
    for line in outline.splitlines():
        l = line.strip()
        if not l:
            continue
        mbul = re.match(r"^[-•–*]\s+(.*)", l)
        mbold = re.match(r"^\*\*(.+?)\*\*:?\s*$", l)
        mnum = re.match(r"^(?:\d+[.)]|#+)\s*(.*)", l)
        if mbul and cur is not None:
            cur[1].append(mbul.group(1).strip())
        elif mbold:
            cur = (mbold.group(1).strip(), []); slides.append(cur)
        elif mnum:
            cur = (mnum.group(1).strip().strip("*"), []); slides.append(cur)
        else:
            if cur is None:
                cur = (l, []); slides.append(cur)
            else:
                cur[1].append(l)
    return slides


@skill("Презентации", "Создать презентацию PowerPoint (.pptx) из ГОТОВОГО плана. outline — план; author — автор. Обычно используй make_slides")
def make_presentation(title, outline, out=None, author=None):
    try:
        from pptx import Presentation
        from pptx.util import Pt
        prs = Presentation()
        # титульный слайд
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = title
        try:
            s.placeholders[1].text = (author or "").strip()
        except Exception:
            pass
        # слайды из плана
        for stitle, bullets in _parse_outline(outline):
            sl = prs.slides.add_slide(prs.slide_layouts[1])
            sl.shapes.title.text = stitle[:120]
            try:
                tf = sl.placeholders[1].text_frame
                tf.clear()
                blist = bullets or [" "]
                for i, b in enumerate(blist):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = b[:200]
                    p.font.size = Pt(20)
            except Exception:
                pass
        out = out or os.path.join(downloads_folder(),
                                  (re.sub(r"\W+", "_", title)[:40] or "presentation") + ".pptx")
        prs.save(out)
        return _ok(f"презентация готова: {out} ({len(prs.slides)} слайдов)")
    except Exception as e:
        return _err(f"презентация: {e}")


_PDF_FONT = None


def _reg_pdf_font():
    """Регистрируем шрифт с кириллицей (Arial из Windows)."""
    global _PDF_FONT
    if _PDF_FONT:
        return _PDF_FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    try:
        pdfmetrics.registerFont(TTFont("UlvarSans", r"C:\Windows\Fonts\arial.ttf"))
        pdfmetrics.registerFont(TTFont("UlvarSans-Bold", r"C:\Windows\Fonts\arialbd.ttf"))
        _PDF_FONT = "UlvarSans"
    except Exception:
        _PDF_FONT = "Helvetica"
    return _PDF_FONT


@skill("Презентации", "Создать презентацию в PDF (слайды 16:9) из ГОТОВОГО плана. outline — заголовки слайдов и пункты; author — кого указать автором. Обычно используй make_slides")
def make_presentation_pdf(title, outline, out=None, author=None):
    try:
        from reportlab.pdfgen import canvas
        from reportlab.lib.colors import HexColor
        font = _reg_pdf_font()
        bold = font + "-Bold" if font == "UlvarSans" else "Helvetica-Bold"
        W, H = 960, 540  # 16:9
        out = out or os.path.join(downloads_folder(),
                                  (re.sub(r"\W+", "_", title)[:40] or "presentation") + ".pdf")
        c = canvas.Canvas(out, pagesize=(W, H))
        accent = HexColor("#d97757")
        dark = HexColor("#1a1a22")
        grey = HexColor("#9a9aa5")
        author = (author or "").strip()

        # титульный слайд
        c.setFillColor(HexColor("#15151c"))
        c.rect(0, 0, W, H, fill=1, stroke=0)
        c.setFillColor(accent)
        c.rect(0, H / 2 + 30, 90, 6, fill=1, stroke=0)
        c.setFillColor(HexColor("#f5f5f7"))
        c.setFont(bold, 46)
        for i, line in enumerate(_wrap(title, 26)[:3]):
            c.drawString(70, H / 2 - 20 - i * 54, line)
        if author:
            c.setFillColor(HexColor("#9a9aa5"))
            c.setFont(font, 18)
            c.drawString(72, 70, author)
        c.showPage()

        slides = _parse_outline(outline)
        for idx, (stitle, bullets) in enumerate(slides, 1):
            c.setFillColor(HexColor("#ffffff"))
            c.rect(0, 0, W, H, fill=1, stroke=0)
            c.setFillColor(accent)
            c.rect(70, H - 96, 70, 6, fill=1, stroke=0)
            c.setFillColor(dark)
            c.setFont(bold, 32)
            for i, line in enumerate(_wrap(stitle, 38)[:2]):
                c.drawString(70, H - 138 - i * 38, line)
            y = H - 200 - (len(_wrap(stitle, 38)[:2]) - 1) * 38
            for b in (bullets or []):
                wrapped = _wrap(b, 64)
                for j, line in enumerate(wrapped):
                    if y < 56:
                        break
                    if j == 0:
                        c.setFillColor(accent)
                        c.circle(84, y + 7, 3.5, fill=1, stroke=0)
                    c.setFillColor(HexColor("#2a2a33"))
                    c.setFont(font, 20)
                    c.drawString(104, y, line)
                    y -= 31
                y -= 7
            # номер слайда (нейтрально, без текста на чужом языке)
            c.setFillColor(grey)
            c.setFont(font, 11)
            c.drawRightString(W - 40, 28, f"{idx} / {len(slides)}")
            c.showPage()
        c.save()
        return _ok(f"презентация в PDF готова: {out} ({len(slides) + 1} слайдов)")
    except Exception as e:
        return _err(f"презентация PDF: {e}")


def _wrap(text, width):
    """Простой перенос по словам до width символов."""
    words, lines, cur = str(text).split(), [], ""
    for w in words:
        if len(cur) + len(w) + 1 <= width:
            cur = (cur + " " + w).strip()
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines or [""]


def _gen_deck(topic, slides=8, language=""):
    """Запасной генератор: заголовок + простой текстовый план (title, outline)."""
    import llm
    import store
    s = store.get_settings()
    lang_line = ("ВАЖНО: весь текст пиши ТОЛЬКО на языке и уровне: " + language + ".\n") if language else ""
    prompt = [
        {"role": "system", "content": "Ты делаешь содержательные презентации. Отвечай строго по формату."},
        {"role": "user", "content": (
            f"Сделай план презентации из {slides} слайдов на тему: «{topic}».\n" + lang_line +
            "Первой строкой: TITLE: <заголовок>\nЗатем слайды:\n1. Заголовок\n- пункт\n- пункт\n- пункт\n")},
    ]
    try:
        raw = llm.complete(prompt, s.get("model"), temperature=0.6, settings=s) or ""
    except Exception:
        raw = ""
    title, body = "", []
    for ln in raw.splitlines():
        m = re.match(r"^\s*TITLE:\s*(.+)$", ln, re.IGNORECASE)
        if m and not title:
            title = m.group(1).strip().strip('"«»')
        else:
            body.append(ln)
    return title, "\n".join(body).strip()


_DECK_SCHEMA = (
    'Каждый слайд — объект с полем "type". Доступные типы и поля:\n'
    '- {"type":"bullets","title":"...","bullets":["...","..."]}  (4-5 пунктов)\n'
    '- {"type":"cards","title":"...","cards":[{"heading":"...","text":"..."},...],"footer":"..."}  (2-3 карточки)\n'
    '- {"type":"compare","title":"...","left":{"heading":"...","text":"..."},"right":{"heading":"...","text":"..."},"footer":"..."}\n'
    '- {"type":"table","title":"...","headers":["...","..."],"rows":[["...","..."],...]}  (2-3 колонки)\n'
    '- {"type":"steps","title":"...","steps":["...","...","..."]}  (3-6 коротких шагов)\n'
    '- {"type":"big","title":"...","big":"КОРОТКО/ЧИСЛО","text":"...","side":{"heading":"...","text":"..."}}  (акцент: большое слово/число)\n'
    '- {"type":"formula","title":"...","rows":[{"label":"...","formula":"...","value":"..."},...],"note":"..."}\n'
    '- {"type":"closing","title":"...","subtitle":"..."}  (финал)\n'
)


def _parse_json_loose(raw):
    """Пытаемся достать JSON из ответа модели, прощая обёртки и мелкие огрехи."""
    import json as _json
    if not raw:
        return None
    s = re.sub(r"^```[a-zA-Z]*", "", raw.strip()).replace("```", "").strip()
    try:
        a, b = s.index("{"), s.rindex("}")
    except ValueError:
        return None
    s = s[a:b + 1]
    # сырые переводы строк/табы внутри JSON (json.loads их не принимает) → пробелы
    s = s.replace("\r", " ").replace("\n", " ").replace("\t", " ")
    s = re.sub(r",\s*([}\]])", r"\1", s).replace(" ", " ")
    try:
        return _json.loads(s)
    except Exception:
        pass
    # вторая попытка: экранируем одиночные обратные слэши
    try:
        return _json.loads(s.replace("\\", "\\\\"))
    except Exception:
        return None


_DECK_EXAMPLE = (
    '{"title":"Gesunder Schlaf","slides":['
    '{"type":"bullets","title":"Warum Schlaf wichtig ist","bullets":["Der Körper erholt sich.",'
    '"Das Gehirn speichert Gelerntes.","Wir sind am Tag konzentrierter."]},'
    '{"type":"cards","title":"Drei Regeln","cards":['
    '{"heading":"Feste Zeiten","text":"Jeden Tag zur gleichen Zeit ins Bett."},'
    '{"heading":"Kein Handy","text":"Eine Stunde vor dem Schlaf kein Bildschirm."},'
    '{"heading":"Dunkel & kühl","text":"Das Zimmer soll dunkel und kühl sein."}]},'
    '{"type":"big","title":"Wie viel Schlaf?","big":"9 Stunden","text":"So viel brauchen Jugendliche."},'
    '{"type":"closing","title":"Schlaf gut!","subtitle":"Gute Nacht."}]}'
)


def _parse_deck_text(raw):
    """Разбирает ответ модели в строчном формате с тегами макетов. (title, [slides]).
    Устойчив к огрехам модели (в отличие от JSON)."""
    title, slides, cur = "", [], None
    for ln in (raw or "").splitlines():
        if not ln.strip():
            continue
        m = re.match(r"^\s*TITLE:\s*(.+)$", ln, re.IGNORECASE)
        if m and not title and cur is None:
            title = m.group(1).strip().strip('"«»'); continue
        mt = re.match(r"^\s*\[(\w+)\]\s*(.*)$", ln)
        if mt:
            typ = mt.group(1).lower()
            cur = {"type": typ, "title": mt.group(2).strip()}
            if typ in ("bullets", "steps"):
                cur["_items"] = []
            elif typ == "cards":
                cur["cards"] = []
            elif typ == "table":
                cur["headers"] = []; cur["rows"] = []
            elif typ == "formula":
                cur["rows"] = []
            slides.append(cur)
            continue
        if cur is None:
            continue
        body, typ = ln.strip(), cur["type"]
        if typ in ("bullets", "steps"):
            if body[:1] in "-•*":
                cur["_items"].append(body.lstrip("-•* ").strip())
        elif typ == "cards":
            if body[:1] in "*-":
                p = body.lstrip("-•* ").split("|", 1)
                cur["cards"].append({"heading": p[0].strip(), "text": (p[1].strip() if len(p) > 1 else "")})
        elif typ == "compare":
            mm = re.match(r"^([LR])\s*:\s*(.*)$", body, re.IGNORECASE)
            if mm:
                p = mm.group(2).split("|", 1)
                cur["left" if mm.group(1).upper() == "L" else "right"] = {
                    "heading": p[0].strip(), "text": (p[1].strip() if len(p) > 1 else "")}
        elif typ == "table":
            if body.startswith("#"):
                cur["headers"] = [x.strip() for x in body.lstrip("# ").split(";")]
            else:
                cur["rows"].append([x.strip() for x in body.split(";")])
        elif typ == "big":
            if body.startswith("="):
                cur["big"] = body.lstrip("= ").strip()
            elif body.startswith("~"):
                cur["text"] = body.lstrip("~ ").strip()
        elif typ == "formula":
            if body.startswith("="):
                p = [x.strip() for x in body.lstrip("= ").split("|")]
                cur["rows"].append({"label": p[0] if p else "",
                                    "formula": p[1] if len(p) > 1 else "",
                                    "value": p[2] if len(p) > 2 else ""})
            elif body.startswith("~"):
                cur["note"] = body.lstrip("~ ").strip()
        elif typ == "closing":
            if body.startswith("~"):
                cur["subtitle"] = body.lstrip("~ ").strip()
    for sl in slides:
        if "_items" in sl:
            sl[("bullets" if sl["type"] == "bullets" else "steps")] = sl.pop("_items")
    return title, [s for s in slides if s.get("type")]


_DECK_FORMAT = (
    "Каждый слайд начинается с тега в [квадратных скобках] и заголовка. Форматы:\n\n"
    "[BULLETS] Заголовок слайда\n- пункт\n- пункт\n- пункт\n\n"
    "[CARDS] Заголовок\n* Подзаголовок | короткий текст\n* Подзаголовок | короткий текст\n\n"
    "[COMPARE] Заголовок\nL: Подзаголовок | текст\nR: Подзаголовок | текст\n\n"
    "[TABLE] Заголовок\n# Колонка1 ; Колонка2 ; Колонка3\nзначение ; значение ; значение\n\n"
    "[STEPS] Заголовок\n- шаг\n- шаг\n\n"
    "[BIG] Заголовок\n= короткое слово или число\n~ пояснение\n\n"
    "[FORMULA] Заголовок\n= что считаем | формула | значение\n\n"
    "[CLOSING] Заголовок\n~ подзаголовок"
)


def _gen_deck_spec(topic, slides=10, language=""):
    """Модель выдаёт структуру слайдов в простом строчном формате с тегами. (title, [slides])."""
    import llm
    import store
    s = store.get_settings()
    lang_line = ("ОЧЕНЬ ВАЖНО: ВЕСЬ текст (заголовки, пункты, подписи) пиши ТОЛЬКО на языке и уровне: "
                 + language + ". Никакого другого языка.\n") if language else ""
    prompt = [
        {"role": "system", "content": "Ты — профессиональный методист и дизайнер презентаций. "
         "Делаешь содержательные слайды с РАЗНЫМИ макетами. Отвечаешь строго в указанном текстовом формате."},
        {"role": "user", "content": (
            f"Сделай презентацию из {slides} слайдов на тему: «{topic}».\n" + lang_line +
            "Первой строкой: TITLE: <общий заголовок>\n"
            "Затем слайды. ЧЕРЕДУЙ макеты: начни с [BULLETS] (введение), дальше используй "
            "[CARDS], [COMPARE], [TABLE], [STEPS], [BIG] (яркий акцент или число), [FORMULA] "
            "(если тема техническая/с расчётами), а последним сделай [CLOSING].\n\n"
            "ТРЕБОВАНИЯ К СОДЕРЖАНИЮ (это важно — не делай пусто):\n"
            "- [BULLETS]: 4-5 пунктов, каждый — законченная фраза с конкретикой.\n"
            "- [CARDS]: РОВНО 3 карточки, в каждой подзаголовок и 1-2 предложения пояснения.\n"
            "- [COMPARE]: в каждой стороне 2-3 предложения.\n"
            "- [TABLE]: 3-4 строки; значения в ячейках КОРОТКИЕ (2-4 слова, не целые предложения).\n"
            "- [STEPS]: 3-5 шагов, подпись каждого — 2-4 слова.\n"
            "- [FORMULA]: 2-4 строки.\n"
            "НЕ пиши в слайдах служебные пометки, числа параметров, кавычки с темой или слова на "
            "другом языке — только содержание по теме. Без markdown-звёздочек.\n\nФОРМАТ:\n" + _DECK_FORMAT)},
    ]
    for attempt in range(2):
        try:
            raw = llm.complete(prompt, s.get("model"), temperature=0.6 if attempt == 0 else 0.3,
                               settings=s, num_predict=4096) or ""
        except Exception:
            raw = ""
        title, sl = _parse_deck_text(raw)
        if len(sl) >= 3:
            return title, sl
    return "", []


def _spec_from_outline(topic, n, language):
    """Запасной план: текстовый outline → слайды-bullets (всё равно через красивый рендерер)."""
    gt, outline = _gen_deck(topic, n, language)
    slides = []
    for stitle, bullets in _parse_outline(outline):
        slides.append({"type": "bullets", "title": stitle, "bullets": bullets or [stitle]})
    if not slides:
        slides = [{"type": "bullets", "title": f"{topic}",
                   "bullets": ["—"]}]
    return gt, slides


_SAN_OK = set("—–…«»„“”€°×·→←↔≈≤≥№%§•’‘\"'()[]{}:;,.!?/\\+-=*&@#№ ")


def _san(t):
    """Чистим текст от символов, которых нет в шрифте (дают «квадраты»), и мусора."""
    out = []
    for ch in str(t):
        if ch in "\n\r\t":
            out.append(" "); continue
        o = ord(ch)
        if o < 32:
            continue
        if o < 0x250 or 0x400 <= o <= 0x4FF or ch in _SAN_OK:  # латиница+расш, кириллица, пунктуация
            out.append(ch)
    return re.sub(r"\s{2,}", " ", "".join(out)).strip()


def _san_spec(obj):
    if isinstance(obj, str):
        return _san(obj)
    if isinstance(obj, list):
        return [_san_spec(x) for x in obj]
    if isinstance(obj, dict):
        return {k: _san_spec(v) for k, v in obj.items()}
    return obj


def _pdf_fonts():
    font = _reg_pdf_font()
    bold = font + "-Bold" if font == "UlvarSans" else "Helvetica-Bold"
    return font, bold


def _render_deck_pdf(title, slides, out, author=""):
    """Рисует презентацию с профессиональными макетами (разные типы слайдов)."""
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor
    font, bold = _pdf_fonts()
    W, H = 960, 540
    NAVY = HexColor("#1b2735"); INK = HexColor("#1f2a37"); ORANGE = HexColor("#ed8b37")
    TEAL = HexColor("#2a9d8f"); CARD = HexColor("#eef1f5"); MUT = HexColor("#5f6b7a")
    WHITE = HexColor("#ffffff"); SUB = HexColor("#9aa5b1"); LINE = HexColor("#dfe3e8")
    author = _san(author or "")
    title = _san(title or "")
    slides = [_san_spec(s) for s in slides if isinstance(s, dict)]
    c = canvas.Canvas(out, pagesize=(W, H))

    def wrp(text, fnt, size, maxw):
        words, lines, cur = str(text).split(), [], ""
        for w in words:
            t = (cur + " " + w).strip()
            if c.stringWidth(t, fnt, size) <= maxw:
                cur = t
            else:
                if cur:
                    lines.append(cur)
                cur = w
        if cur:
            lines.append(cur)
        return lines or [""]

    def header(stitle):
        c.setFillColor(WHITE); c.rect(0, 0, W, H, fill=1, stroke=0)
        c.setFillColor(ORANGE); c.rect(64, H - 72, 52, 5, fill=1, stroke=0)
        c.setFillColor(INK); c.setFont(bold, 30)
        for i, ln in enumerate(wrp(stitle, bold, 30, W - 128)[:2]):
            c.drawString(64, H - 104 - i * 34, ln)

    def pagenum(idx):
        c.setFillColor(SUB); c.setFont(font, 10)
        c.drawRightString(W - 40, 26, str(idx))

    def badge(cx, cy, rad, color, txt):
        c.setFillColor(color); c.circle(cx, cy, rad, fill=1, stroke=0)
        c.setFillColor(WHITE); c.setFont(bold, rad * 1.05)
        c.drawCentredString(cx, cy - rad * 0.36, str(txt))

    def para(text, x, y, fnt, size, color, maxw, lh, maxlines=20):
        c.setFont(fnt, size); c.setFillColor(color)
        for ln in wrp(text, fnt, size, maxw)[:maxlines]:
            c.drawString(x, y, ln); y -= lh
        return y

    # ---- титульный слайд ----
    c.setFillColor(NAVY); c.rect(0, 0, W, H, fill=1, stroke=0)
    tlines = wrp(title or "Презентация", bold, 46, W - 200)[:3]
    top = H * 0.52 + (len(tlines) - 1) * 28      # верх блока заголовка
    c.setFillColor(ORANGE); c.rect(64, top + 34, 90, 6, fill=1, stroke=0)   # полоса НАД заголовком
    c.setFillColor(WHITE); c.setFont(bold, 46)
    ty = top
    for ln in tlines:
        c.drawString(64, ty, ln); ty -= 56
    if author:
        c.setFillColor(SUB); c.setFont(font, 17); c.drawString(66, 70, author)
    c.showPage()

    total = len(slides)
    for i, sl in enumerate(slides, 1):
        if not isinstance(sl, dict):
            continue
        typ = str(sl.get("type", "bullets")).lower()
        st = sl.get("title", "")
        try:
            if typ == "closing":
                c.setFillColor(ORANGE); c.rect(0, 0, W, H, fill=1, stroke=0)
                c.setFillColor(WHITE); c.setFont(bold, 44)
                cl = wrp(st or sl.get("big", ""), bold, 44, W - 160)[:3]
                yy = H / 2 + (len(cl) - 1) * 26 + 10
                for ln in cl:
                    c.drawCentredString(W / 2, yy, ln); yy -= 52
                sub = sl.get("subtitle", "")
                if sub:
                    c.setFont(font, 20)
                    yy -= 6
                    for ln in wrp(sub, font, 20, W - 220)[:2]:
                        c.drawCentredString(W / 2, yy, ln); yy -= 28

            elif typ == "big":
                c.setFillColor(NAVY); c.rect(0, 0, W, H, fill=1, stroke=0)
                c.setFillColor(WHITE); c.setFont(bold, 22)
                c.drawString(64, H - 86, st[:70])
                side = sl.get("side") or {}
                bigw = (W * 0.52) if side else (W - 160)
                big = str(sl.get("big", "")).strip()
                bs = 92
                while bs > 40 and c.stringWidth(big, bold, bs) > bigw:
                    bs -= 4
                c.setFillColor(ORANGE); c.setFont(bold, bs)
                c.drawString(64, H * 0.5, big)
                yy = H * 0.5 - 30
                if sl.get("text"):
                    yy = para(sl["text"], 66, yy, font, 18, SUB, bigw, 24, 4)
                if side:
                    px, pw = W * 0.60, W * 0.34
                    c.setFillColor(HexColor("#24303f"))
                    c.roundRect(px, 90, pw, H - 200, 14, fill=1, stroke=0)
                    yy2 = H - 150
                    c.setFillColor(TEAL); c.setFont(bold, 18)
                    for ln in wrp(side.get("heading", ""), bold, 18, pw - 44)[:2]:
                        c.drawString(px + 22, yy2, ln); yy2 -= 24
                    para(side.get("text", ""), px + 22, yy2 - 6, font, 15, HexColor("#cfd6de"), pw - 44, 22, 8)
                pagenum(i)

            elif typ == "cards":
                header(st)
                cards = [x for x in (sl.get("cards") or []) if isinstance(x, dict)][:3] or [{}]
                n = len(cards); gap = 18; x0 = 64
                cw = (W - 128 - (n - 1) * gap) / n
                cy, ch = 90, 260
                for k, cd in enumerate(cards):
                    x = x0 + k * (cw + gap)
                    c.setFillColor(CARD); c.roundRect(x, cy, cw, ch, 14, fill=1, stroke=0)
                    head = wrp(cd.get("heading", ""), bold, 17, cw - 44)[:2]
                    body = wrp(cd.get("text", ""), font, 14, cw - 44)[:10]
                    content_h = 34 + 14 + len(head) * 22 + 6 + len(body) * 20
                    toppad = max(26, (ch - content_h) / 2)
                    yb = cy + ch - toppad           # верх контента (центрирован)
                    badge(x + 34, yb - 16, 17, ORANGE if k % 2 == 0 else TEAL, k + 1)
                    yy = yb - 56
                    c.setFillColor(INK); c.setFont(bold, 17)
                    for ln in head:
                        c.drawString(x + 22, yy, ln); yy -= 22
                    yy -= 6
                    c.setFillColor(MUT); c.setFont(font, 14)
                    for ln in body:
                        c.drawString(x + 22, yy, ln); yy -= 20
                if sl.get("footer"):
                    c.setFillColor(MUT); c.setFont(font, 13)
                    c.drawCentredString(W / 2, 50, str(sl["footer"])[:130])
                pagenum(i)

            elif typ == "compare":
                header(st)
                gap = 18; pw = (W - 128 - gap) / 2; py, ph = 90, 260
                for k, key in enumerate(("left", "right")):
                    d = sl.get(key) or {}
                    x = 64 + k * (pw + gap)
                    c.setFillColor(CARD); c.roundRect(x, py, pw, ph, 14, fill=1, stroke=0)
                    head = str(d.get("heading", ""))[:40]
                    body = wrp(d.get("text", ""), font, 15, pw - 48)[:10]
                    content_h = 30 + 16 + len(body) * 21
                    toppad = max(30, (ph - content_h) / 2)
                    yb = py + ph - toppad
                    c.setFillColor(ORANGE if k == 0 else TEAL); c.setFont(bold, 21)
                    c.drawString(x + 24, yb, head)
                    yy = yb - 34
                    c.setFillColor(INK); c.setFont(font, 15)
                    for ln in body:
                        c.drawString(x + 24, yy, ln); yy -= 21
                if sl.get("footer"):
                    c.setFillColor(MUT); c.setFont(font, 13)
                    c.drawCentredString(W / 2, 50, str(sl["footer"])[:130])
                pagenum(i)

            elif typ == "table":
                header(st)
                headers = [str(h) for h in (sl.get("headers") or [])][:4]
                rows = [r for r in (sl.get("rows") or []) if isinstance(r, list)][:7]
                ncol = max(1, len(headers) or (len(rows[0]) if rows else 1))
                x0, tw = 64, W - 128
                colw = tw / ncol
                topY, botY = H - 162, 64
                fs = 12.5
                # высота строки под перенос (до 2 строк)
                def cell_lines(txt):
                    return wrp(txt, font, fs, colw - 24)[:2]
                hH = 40
                rowHs = []
                for r in rows:
                    mx = 1
                    for j in range(ncol):
                        mx = max(mx, len(cell_lines(str(r[j]) if j < len(r) else "")))
                    rowHs.append(22 + mx * 16)
                need = hH + sum(rowHs)
                avail = topY - botY
                scale = min(1.0, avail / need) if need > avail else 1.0
                hH *= scale; rowHs = [h * scale for h in rowHs]; fs *= max(0.8, scale)
                y = topY - hH
                if headers:
                    c.setFillColor(NAVY); c.rect(x0, y, tw, hH, fill=1, stroke=0)
                    c.setFillColor(WHITE); c.setFont(bold, fs + 0.5)
                    for j, h in enumerate(headers):
                        c.drawString(x0 + j * colw + 12, y + hH / 2 - 5, h[:30])
                for ri, r in enumerate(rows):
                    rh = rowHs[ri]; y -= rh
                    c.setStrokeColor(LINE); c.setLineWidth(1); c.line(x0, y, x0 + tw, y)
                    for j in range(ncol):
                        cell = str(r[j]) if j < len(r) else ""
                        c.setFillColor(INK if j == 0 else MUT)
                        c.setFont(bold if j == 0 else font, fs)
                        lines = wrp(cell, bold if j == 0 else font, fs, colw - 24)[:2]
                        ty2 = y + rh - 18
                        for ln in lines:
                            c.drawString(x0 + j * colw + 12, ty2, ln); ty2 -= 16
                pagenum(i)

            elif typ == "steps":
                header(st)
                steps = [str(x)[:70] for x in (sl.get("steps") or [])][:5] or [""]
                n = len(steps); cy = H * 0.52; rad = 26
                x0 = 64 + (W - 128) / (2 * n)
                span = (W - 128) / n
                prev = None
                for k, lab in enumerate(steps):
                    cx = x0 + k * span
                    if prev is not None:
                        c.setStrokeColor(HexColor("#9aa5b1")); c.setLineWidth(2)
                        c.line(prev + rad + 8, cy, cx - rad - 8, cy)
                    badge(cx, cy, rad, ORANGE if k % 2 == 0 else TEAL, k + 1)
                    c.setFillColor(INK); c.setFont(bold, 13)
                    yy = cy - rad - 22
                    for ln in wrp(lab, bold, 13, span - 24)[:3]:
                        c.drawCentredString(cx, yy, ln); yy -= 16
                    prev = cx
                pagenum(i)

            elif typ == "formula":
                header(st)
                rows = [r for r in (sl.get("rows") or []) if isinstance(r, dict)][:6] or [{}]
                topY = H - 170; barH = min(54, (topY - 90) / len(rows) - 10); gap = 12
                y = topY - barH
                for k, r in enumerate(rows):
                    last = (k == len(rows) - 1)
                    c.setFillColor(NAVY if last else CARD)
                    c.roundRect(64, y, W - 128, barH, 10, fill=1, stroke=0)
                    c.setFillColor(WHITE if last else INK); c.setFont(bold, 15)
                    c.drawString(84, y + barH / 2 - 6, str(r.get("label", ""))[:34])
                    c.setFillColor(ORANGE if last else TEAL); c.setFont(bold, 17)
                    c.drawString(64 + (W - 128) * 0.36, y + barH / 2 - 6, str(r.get("formula", ""))[:40])
                    c.setFillColor(HexColor("#dfe3e8") if last else MUT); c.setFont(font, 13)
                    c.drawRightString(W - 86, y + barH / 2 - 5, str(r.get("value", ""))[:40])
                    y -= barH + gap
                if sl.get("note"):
                    c.setFillColor(MUT); c.setFont(font, 13)
                    c.drawCentredString(W / 2, 52, sl["note"][:130])
                pagenum(i)

            else:  # bullets (по умолчанию)
                header(st)
                y = H - 176
                for b in [str(x) for x in (sl.get("bullets") or [])][:7]:
                    lines = wrp(b, font, 19, W - 220)
                    c.setFillColor(ORANGE); c.circle(82, y + 6, 3.5, fill=1, stroke=0)
                    c.setFillColor(HexColor("#2a3340")); c.setFont(font, 19)
                    for j, ln in enumerate(lines):
                        c.drawString(102, y, ln); y -= 27
                    y -= 10
                pagenum(i)
        except Exception:
            pagenum(i)
        c.showPage()
    c.save()
    return _ok(f"презентация в PDF готова: {out} ({total + 1} слайдов)")


@skill("Презентации", "Сделать простую аккуратную презентацию по теме: Ульвар САМ придумает содержание (слайды «заголовок + пункты») и заголовок на нужном языке. Параметры: topic (тема), slides (число слайдов, по умолчанию 8), language (язык и уровень, напр. 'немецкий A2' или 'русский'), pdf (True=PDF, False=PowerPoint), author (автор на титуле). Используй для ЛЮБОЙ просьбы сделать презентацию.")
def make_slides(topic, slides=8, language="", pdf=True, author=None):
    try:
        n = max(4, min(int(slides or 8), 14))
    except Exception:
        n = 8
    if author is None:
        import store
        author = store.get_settings().get("user_name", "") or ""
    use_pdf = pdf if isinstance(pdf, bool) else str(pdf).lower() not in ("false", "0", "pptx", "нет")

    if use_pdf:
        # простые, чистые слайды «заголовок + пункты» — надёжно и аккуратно
        gen_title, spec = _spec_from_outline(topic, n, language)
        title = (gen_title or str(topic)).strip()[:80] or "Презентация"
        out = os.path.join(downloads_folder(), (re.sub(r"\W+", "_", title)[:40] or "presentation") + ".pdf")
        return _render_deck_pdf(title, spec, out, author=author)

    # PowerPoint
    gt, outline = _gen_deck(topic, n, language)
    title = (gt or str(topic)).strip()[:80] or "Презентация"
    if not _parse_outline(outline):
        outline = "\n".join(f"{k}. {topic} — {k}\n- Пункт\n- Детали\n- Пример" for k in range(1, n + 1))
    return make_presentation(title, outline, author=author)


@skill("Презентации", "Сколько слайдов в .pptx")
def pptx_slides_count(path):
    try:
        from pptx import Presentation
        return f"Слайдов: {len(Presentation(path).slides)}"
    except Exception as e:
        return _err(str(e))


@skill("Презентации", "Извлечь текст из презентации .pptx")
def pptx_text(path):
    try:
        from pptx import Presentation
        out = []
        for i, sl in enumerate(Presentation(path).slides, 1):
            texts = [sh.text for sh in sl.shapes if getattr(sh, "has_text_frame", False) and sh.text]
            out.append(f"Слайд {i}: " + " | ".join(texts))
        return "\n".join(out)
    except Exception as e:
        return _err(str(e))
