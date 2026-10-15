# -*- coding: utf-8 -*-
"""Навыки Ульвара: документы (PDF, Excel, Word)."""
from skills import skill, _ok, _err
import os


# ---------- PDF ----------
@skill("PDF", "Извлечь текст из PDF")
def pdf_text(path, max_chars=20000):
    try:
        from pypdf import PdfReader
        r = PdfReader(path)
        t = "\n".join((p.extract_text() or "") for p in r.pages)
        return t[:int(max_chars)] or "(текст не извлёкся — возможно скан)"
    except Exception as e:
        return _err(str(e))


@skill("PDF", "Сколько страниц в PDF")
def pdf_pages(path):
    try:
        from pypdf import PdfReader
        return f"Страниц: {len(PdfReader(path).pages)}"
    except Exception as e:
        return _err(str(e))


@skill("PDF", "Объединить несколько PDF в один")
def pdf_merge(paths, out="merged.pdf"):
    try:
        from pypdf import PdfWriter
        if isinstance(paths, str):
            paths = [p.strip() for p in paths.split(";")]
        w = PdfWriter()
        for p in paths:
            w.append(p)
        w.write(out); w.close()
        return _ok(f"объединено → {out}")
    except Exception as e:
        return _err(str(e))


@skill("PDF", "Разбить PDF на отдельные страницы")
def pdf_split(path, out_dir=None):
    try:
        from pypdf import PdfReader, PdfWriter
        out_dir = out_dir or os.path.splitext(path)[0] + "_pages"
        os.makedirs(out_dir, exist_ok=True)
        r = PdfReader(path)
        for i, page in enumerate(r.pages, 1):
            w = PdfWriter(); w.add_page(page)
            with open(os.path.join(out_dir, f"page_{i}.pdf"), "wb") as f:
                w.write(f)
        return _ok(f"{len(r.pages)} страниц → {out_dir}")
    except Exception as e:
        return _err(str(e))


@skill("PDF", "Вытащить диапазон страниц в новый PDF")
def pdf_extract_pages(path, start, end, out="extract.pdf"):
    try:
        from pypdf import PdfReader, PdfWriter
        r = PdfReader(path); w = PdfWriter()
        for i in range(int(start)-1, int(end)):
            w.add_page(r.pages[i])
        with open(out, "wb") as f:
            w.write(f)
        return _ok(f"страницы {start}-{end} → {out}")
    except Exception as e:
        return _err(str(e))


@skill("PDF", "Повернуть все страницы PDF")
def pdf_rotate(path, degrees=90, out=None):
    try:
        from pypdf import PdfReader, PdfWriter
        r = PdfReader(path); w = PdfWriter()
        for p in r.pages:
            p.rotate(int(degrees)); w.add_page(p)
        out = out or os.path.splitext(path)[0] + "_rot.pdf"
        with open(out, "wb") as f:
            w.write(f)
        return _ok(out)
    except Exception as e:
        return _err(str(e))


@skill("PDF", "Поставить пароль на PDF")
def pdf_encrypt(path, password, out=None):
    try:
        from pypdf import PdfReader, PdfWriter
        r = PdfReader(path); w = PdfWriter()
        for p in r.pages:
            w.add_page(p)
        w.encrypt(password)
        out = out or os.path.splitext(path)[0] + "_protected.pdf"
        with open(out, "wb") as f:
            w.write(f)
        return _ok(out)
    except Exception as e:
        return _err(str(e))


@skill("PDF", "Картинки → PDF")
def images_to_pdf(paths, out="images.pdf"):
    try:
        from PIL import Image
        if isinstance(paths, str):
            paths = [p.strip() for p in paths.split(";")]
        ims = [Image.open(p).convert("RGB") for p in paths]
        ims[0].save(out, save_all=True, append_images=ims[1:])
        return _ok(f"{len(ims)} картинок → {out}")
    except Exception as e:
        return _err(str(e))


@skill("PDF", "PDF-страницы → картинки (через PyMuPDF, если есть)")
def pdf_to_images(path, out_dir=None):
    try:
        import fitz  # PyMuPDF
        out_dir = out_dir or os.path.splitext(path)[0] + "_img"
        os.makedirs(out_dir, exist_ok=True)
        doc = fitz.open(path)
        for i, page in enumerate(doc, 1):
            page.get_pixmap(dpi=150).save(os.path.join(out_dir, f"page_{i}.png"))
        return _ok(f"{len(doc)} картинок → {out_dir}")
    except ImportError:
        return _err("нужен PyMuPDF: pip install pymupdf")
    except Exception as e:
        return _err(str(e))


# ---------- Excel ----------
@skill("Excel", "Прочитать Excel и показать первые строки")
def excel_preview(path, rows=10):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        ws = wb.active
        out = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= int(rows):
                break
            out.append(", ".join(str(c) if c is not None else "" for c in row))
        return "\n".join(out)
    except Exception as e:
        return _err(str(e))


@skill("Excel", "Размер Excel-таблицы (строк × колонок)")
def excel_shape(path):
    try:
        import openpyxl
        ws = openpyxl.load_workbook(path, read_only=True).active
        return f"Строк: {ws.max_row}, колонок: {ws.max_column}"
    except Exception as e:
        return _err(str(e))


@skill("Excel", "Список листов в книге Excel")
def excel_sheets(path):
    try:
        import openpyxl
        return ", ".join(openpyxl.load_workbook(path, read_only=True).sheetnames)
    except Exception as e:
        return _err(str(e))


@skill("Excel", "Excel → CSV")
def excel_to_csv(path, out=None):
    try:
        import openpyxl, csv
        ws = openpyxl.load_workbook(path, read_only=True).active
        out = out or os.path.splitext(path)[0] + ".csv"
        with open(out, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                w.writerow(row)
        return _ok(out)
    except Exception as e:
        return _err(str(e))


@skill("Excel", "Создать Excel из CSV")
def csv_to_excel(path, out=None):
    try:
        import openpyxl, csv
        wb = openpyxl.Workbook(); ws = wb.active
        with open(path, encoding="utf-8", errors="replace") as f:
            for row in csv.reader(f):
                ws.append(row)
        out = out or os.path.splitext(path)[0] + ".xlsx"
        wb.save(out); return _ok(out)
    except Exception as e:
        return _err(str(e))


@skill("Excel", "Сумма числовой колонки в Excel (по букве, напр. B)")
def excel_sum_column(path, column="A"):
    try:
        import openpyxl
        ws = openpyxl.load_workbook(path, read_only=True).active
        total = 0.0
        for cell in ws[column]:
            if isinstance(cell.value, (int, float)):
                total += cell.value
        return f"Сумма колонки {column}: {total:g}"
    except Exception as e:
        return _err(str(e))


@skill("Excel", "Создать Excel из таблицы (строки через ;, ячейки через ,)")
def make_excel(rows_text, out="table.xlsx"):
    try:
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        for line in rows_text.split(";"):
            ws.append([c.strip() for c in line.split(",")])
        wb.save(out); return _ok(out)
    except Exception as e:
        return _err(str(e))


# ---------- Word ----------
@skill("Word", "Прочитать текст из Word-документа")
def word_text(path, max_chars=20000):
    try:
        import docx
        d = docx.Document(path)
        return "\n".join(p.text for p in d.paragraphs)[:int(max_chars)]
    except Exception as e:
        return _err(str(e))


@skill("Word", "Создать Word-документ с текстом")
def make_word(text, out="document.docx"):
    try:
        import docx
        d = docx.Document()
        for line in text.split("\n"):
            d.add_paragraph(line)
        d.save(out); return _ok(out)
    except Exception as e:
        return _err(str(e))


@skill("Word", "Word → PDF (через Word, если установлен)")
def word_to_pdf(path, out=None):
    try:
        from docx2pdf import convert
        out = out or os.path.splitext(path)[0] + ".pdf"
        convert(path, out); return _ok(out)
    except ImportError:
        return _err("нужен docx2pdf и установленный Word")
    except Exception as e:
        return _err(str(e))


@skill("Word", "Сколько слов в Word-документе")
def word_count_doc(path):
    try:
        import docx
        d = docx.Document(path)
        words = sum(len(p.text.split()) for p in d.paragraphs)
        return f"Слов: {words}, абзацев: {len(d.paragraphs)}"
    except Exception as e:
        return _err(str(e))


@skill("Word", "Добавить абзац в существующий Word-документ")
def word_append(path, text):
    try:
        import docx
        d = docx.Document(path)
        d.add_paragraph(text); d.save(path)
        return _ok("добавлено")
    except Exception as e:
        return _err(str(e))


@skill("Документы", "Текстовый файл → PDF")
def text_to_pdf(path, out=None):
    try:
        from PIL import Image, ImageDraw, ImageFont
        with open(path, encoding="utf-8", errors="replace") as f:
            lines = f.read().split("\n")
        out = out or os.path.splitext(path)[0] + ".pdf"
        pages = []
        font = ImageFont.load_default()
        for i in range(0, len(lines), 50):
            img = Image.new("RGB", (800, 1100), "white")
            d = ImageDraw.Draw(img)
            for j, line in enumerate(lines[i:i+50]):
                d.text((40, 30+j*20), line[:110], fill="black", font=font)
            pages.append(img)
        if pages:
            pages[0].save(out, save_all=True, append_images=pages[1:])
        return _ok(out)
    except Exception as e:
        return _err(str(e))
